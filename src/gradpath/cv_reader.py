"""Read text from PDF and DOCX CV files."""

from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


MIN_TEXT_CHARACTERS = 40


def clean_extracted_text(text: str) -> str:
    """Normalize whitespace while preserving meaningful line breaks."""
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]*\n+", "\n\n", text)
    return text.strip()


def read_pdf(path: Path) -> tuple[str, list[str]]:
    """Extract text from a PDF and report extraction warnings."""
    warnings = []
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
            pages.append(page_text)
        except Exception as exc:
            pages.append("")
            warnings.append(
                f"Could not extract text from PDF page {page_number}: "
                f"{type(exc).__name__}"
            )

    return "\n\n".join(pages), warnings


def read_docx(path: Path) -> tuple[str, list[str]]:
    """Extract text from DOCX paragraphs and tables."""
    document = Document(str(path))
    sections = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            sections.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                sections.append(" | ".join(cells))

    return "\n".join(sections), []


def read_cv(file_path: str | Path) -> dict:
    """
    Read a PDF or DOCX CV and return text plus extraction metadata.

    The returned object does not include the original file path or
    the original document contents beyond the extracted text itself.
    Callers should avoid persisting extracted_text unnecessarily.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"CV file not found: {path.name}")

    if not path.is_file():
        raise ValueError("The provided CV path is not a file.")

    extension = path.suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise ValueError(
            "Unsupported CV format. Please provide a .pdf or .docx file."
        )

    if path.stat().st_size == 0:
        raise ValueError("The CV file is empty.")

    if extension == ".pdf":
        raw_text, warnings = read_pdf(path)
    else:
        raw_text, warnings = read_docx(path)

    text = clean_extracted_text(raw_text)

    if len(text) < MIN_TEXT_CHARACTERS:
        warnings.append(
            "Very little text was extracted. The document may be scanned, "
            "image-based, empty, or otherwise difficult to parse."
        )

    status = "failed" if not text else (
        "partial" if warnings else "success"
    )

    return {
        "file_type": extension.lstrip("."),
        "extraction_status": status,
        "text_char_count": len(text),
        "extracted_text": text,
        "extraction_warnings": warnings,
    }
