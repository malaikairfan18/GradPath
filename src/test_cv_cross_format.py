
"""Automated cross-format tests for equivalent PDF and DOCX CVs."""

from pathlib import Path

import pytest
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from gradpath.cv_profile import build_candidate_profile
from gradpath.profile_validator import validate_candidate_profile


CV_TEXT = [
    ("Technical Skills", "Python, SQL, Power BI"),
    ("Education", "BS Data Science, Example University"),
    ("Projects", "Built a dashboard using Python and SQL."),
]


def create_docx(path: Path) -> None:
    doc = Document()

    for heading, content in CV_TEXT:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(content)

    doc.save(path)


def create_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    y = 750

    for heading, content in CV_TEXT:
        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(60, y, heading)
        y -= 22

        pdf.setFont("Helvetica", 10)
        pdf.drawString(60, y, content)
        y -= 35

    pdf.save()


@pytest.fixture
def profile_pair(tmp_path):
    """Build profiles from equivalent PDF and DOCX sample CVs."""
    docx_path = tmp_path / "sample.docx"
    pdf_path = tmp_path / "sample.pdf"

    create_docx(docx_path)
    create_pdf(pdf_path)

    return (
        build_candidate_profile(docx_path),
        build_candidate_profile(pdf_path),
    )


def test_both_formats_extract_successfully(profile_pair):
    docx_profile, pdf_profile = profile_pair

    assert docx_profile["source_document"]["extraction_status"] == "success"
    assert pdf_profile["source_document"]["extraction_status"] == "success"


def test_equivalent_formats_detect_the_same_skills(profile_pair):
    docx_profile, pdf_profile = profile_pair

    docx_skills = {
        skill["skill_name"] for skill in docx_profile["skills"]
    }
    pdf_skills = {
        skill["skill_name"] for skill in pdf_profile["skills"]
    }

    assert docx_skills == pdf_skills
    assert docx_skills == {"Python", "SQL", "Power BI"}


def test_both_profiles_pass_validation(profile_pair):
    docx_profile, pdf_profile = profile_pair

    assert validate_candidate_profile(docx_profile) == []
    assert validate_candidate_profile(pdf_profile) == []


def test_profiles_exclude_raw_extracted_text(profile_pair):
    docx_profile, pdf_profile = profile_pair

    assert "extracted_text" not in docx_profile
    assert "extracted_text" not in pdf_profile


def test_extracted_skills_remain_unverified(profile_pair):
    docx_profile, pdf_profile = profile_pair

    for profile in (docx_profile, pdf_profile):
        assert profile["skills"]
        assert all(
            skill["verification_status"] == "unverified"
            for skill in profile["skills"]
        )
