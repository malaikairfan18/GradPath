from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from gradpath.cv_profile import build_candidate_profile


def build_test_cv(path: Path) -> dict:
    document = Document()

    document.add_heading("Education", level=1)
    document.add_paragraph("BS Data Science")
    document.add_paragraph("GIFT University, Gujranwala")
    document.add_paragraph("2022 - 2026")
    document.add_paragraph(
        "Relevant Coursework: Machine Learning, Statistics"
    )

    document.add_heading("Experience", level=1)
    document.add_paragraph("Data Analyst Intern")
    document.add_paragraph("Example Company | June 2025 - August 2025")
    document.add_paragraph("Cleaned datasets using Python and SQL.")
    document.add_paragraph("Created dashboards to track business KPIs.")

    document.add_heading("Projects", level=1)
    document.add_paragraph("Sales Analytics Dashboard")
    document.add_paragraph("Tools: Python, SQL, Power BI")
    document.add_paragraph("Analyzed sales trends and visualized KPIs.")
    document.add_paragraph("Fake News Detection")
    document.add_paragraph(
        "Built a text classification model using machine learning."
    )

    document.save(path)
    return build_candidate_profile(path)


def test_education_details_are_grouped_into_one_record():
    with TemporaryDirectory() as folder:
        profile = build_test_cv(Path(folder) / "cv.docx")

    assert len(profile["education"]) == 1
    education = profile["education"][0]
    combined = " ".join(str(value) for value in education.values())

    assert "BS Data Science" in combined
    assert "GIFT University" in combined
    assert "2022 - 2026" in combined
    assert "Machine Learning" in combined
    assert "Statistics" in combined


def test_experience_details_are_grouped_into_one_record():
    with TemporaryDirectory() as folder:
        profile = build_test_cv(Path(folder) / "cv.docx")

    assert len(profile["experience"]) == 1
    experience = profile["experience"][0]
    combined = " ".join(str(value) for value in experience.values())

    assert "Data Analyst Intern" in combined
    assert "Example Company" in combined
    assert "June 2025" in combined
    assert "August 2025" in combined
    assert "Cleaned datasets" in combined
    assert "business KPIs" in combined


def test_projects_are_separated_and_keep_their_details():
    with TemporaryDirectory() as folder:
        profile = build_test_cv(Path(folder) / "cv.docx")

    assert len(profile["projects"]) == 2

    first, second = profile["projects"]
    assert first["name"] == "Sales Analytics Dashboard"
    assert "Analyzed sales trends" in first["description"]
    assert {"Python", "SQL", "Power BI"}.issubset(
        set(first["technologies"])
    )

    assert second["name"] == "Fake News Detection"
    assert "text classification model" in second["description"]


def test_profile_does_not_expose_full_cv_text():
    with TemporaryDirectory() as folder:
        profile = build_test_cv(Path(folder) / "cv.docx")

    assert "extracted_text" not in profile
    assert all(
        skill["verification_status"] == "unverified"
        for skill in profile["skills"]
    )


def test_multiple_degrees_are_kept_as_separate_education_records():
    from gradpath.cv_profile import extract_entries

    education = extract_entries(
        "BS Data Science\nGIFT University\n"
        "MS Computer Science\nExample Institute",
        "education",
    )

    assert len(education) == 2
    assert "BS Data Science" in education[0]["details"]
    assert "GIFT University" in education[0]["details"]
    assert "MS Computer Science" in education[1]["details"]
    assert "Example Institute" in education[1]["details"]


def test_experience_responsibility_is_not_mistaken_for_a_new_role():
    from gradpath.cv_profile import extract_entries

    experience = extract_entries(
        "Data Analyst Intern\n"
        "Example Company | June 2025 - August 2025\n"
        "Supported the data analyst with reporting.\n"
        "Created dashboards for business KPIs.",
        "experience",
    )

    assert len(experience) == 1
    assert experience[0]["job_title"] == "Data Analyst Intern"
    assert "Supported the data analyst with reporting." in experience[0]["description"]
    assert "Created dashboards for business KPIs." in experience[0]["description"]


def test_sentence_case_project_title_is_not_lost_in_description():
    from gradpath.cv_profile import extract_entries

    projects = extract_entries(
        "Sales analytics dashboard\n"
        "Analyzed sales trends and visualized KPIs.\n"
        "Fake news detection system\n"
        "Built a text classification model.",
        "projects",
    )

    assert len(projects) == 2
    assert projects[0]["name"] == "Sales analytics dashboard"
    assert "Analyzed sales trends" in projects[0]["description"]
    assert projects[1]["name"] == "Fake news detection system"
    assert "text classification model" in projects[1]["description"]
