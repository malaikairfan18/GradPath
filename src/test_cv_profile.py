from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from gradpath.cv_profile import build_candidate_profile


def main():
    with TemporaryDirectory() as temp_dir:
        cv_path = Path(temp_dir) / "synthetic_cv.docx"

        document = Document()
        document.add_heading("Alex Example", level=1)
        document.add_heading("Education", level=2)
        document.add_paragraph("BS Data Science, Example University")
        document.add_heading("Technical Skills", level=2)
        document.add_paragraph("Python, SQL, Power BI")
        document.add_heading("Projects", level=2)
        document.add_paragraph(
            "Built a sales dashboard using Python and SQL."
        )
        document.save(cv_path)

        profile = build_candidate_profile(
            cv_path,
            candidate_id="test_candidate_001",
        )

        print("File type:", profile["source_document"]["file_type"])
        print("Extraction status:", profile["source_document"]["extraction_status"])
        print("Education entries:", len(profile["education"]))
        print("Project entries:", len(profile["projects"]))
        print("Extracted skills:", [s["skill_name"] for s in profile["skills"]])
        print("Full CV text excluded:", "extracted_text" not in profile)
        print(
            "Skills remain unverified:",
            all(
                s["verification_status"] == "unverified"
                for s in profile["skills"]
            ),
        )


if __name__ == "__main__":
    main()
