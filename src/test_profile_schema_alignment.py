from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from gradpath.cv_profile import build_candidate_profile


with TemporaryDirectory() as folder:
    path = Path(folder) / "sample.docx"

    doc = Document()
    doc.add_heading("Education", level=1)
    doc.add_paragraph("BS Data Science, Example University")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Built a dashboard using Python.")
    doc.save(path)

    profile = build_candidate_profile(path)

    print("Education schema fields:", sorted(profile["education"][0].keys()))
    print("Project schema fields:", sorted(profile["projects"][0].keys()))
    print(
        "Unknown degree remains null:",
        profile["education"][0]["degree"] is None,
    )
    print(
        "Project description retained:",
        "Python" in profile["projects"][0]["description"],
    )
