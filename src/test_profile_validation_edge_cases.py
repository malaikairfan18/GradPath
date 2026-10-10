from pathlib import Path
from tempfile import TemporaryDirectory
import copy

from docx import Document

from gradpath.cv_profile import build_candidate_profile
from gradpath.profile_validator import validate_candidate_profile


with TemporaryDirectory() as folder:
    path = Path(folder) / "sample.docx"

    doc = Document()
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, SQL")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Used Python and SQL to analyze data.")
    doc.save(path)

    profile = build_candidate_profile(path)

    if not profile["skills"]:
        raise AssertionError("The test CV should produce at least one skill.")

    # Test 1: Empty evidence text must be rejected.
    empty_evidence = copy.deepcopy(profile)
    empty_evidence["skills"][0]["evidence"][0]["text"] = ""
    errors = validate_candidate_profile(empty_evidence)
    print(
        "1. Empty evidence rejected:",
        any("non-empty text" in error for error in errors),
    )

    # Test 2: Invalid verification status must be rejected.
    invalid_status = copy.deepcopy(profile)
    invalid_status["skills"][0]["verification_status"] = "probably_verified"
    errors = validate_candidate_profile(invalid_status)
    print(
        "2. Invalid verification status rejected:",
        any("invalid verification_status" in error for error in errors),
    )

    # Test 3: CV extraction must not auto-verify skills.
    print(
        "3. Extracted skills remain unverified:",
        all(
            skill["verification_status"] == "unverified"
            for skill in profile["skills"]
        ),
    )

    # Test 4: The generated profile must not contain raw CV text.
    print(
        "4. Raw CV text excluded:",
        "extracted_text" not in profile,
    )

    # Test 5: A valid profile should still pass.
    print(
        "5. Valid profile passes:",
        validate_candidate_profile(profile) == [],
    )
